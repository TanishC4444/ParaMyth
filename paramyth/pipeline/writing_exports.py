"""Export the author's current chapter versions in their selected medium."""
import io
import re
from pathlib import Path
from .storage import load
from .creative_studio import scene_output
from .media import contract


def manuscript(folder,scene_id=''):
    author=load(folder/'author-workspace.json') if (folder/'author-workspace.json').is_file() else {}
    drafts=load(folder/'workspace-drafts.json') if (folder/'workspace-drafts.json').is_file() else {}
    scenes=list(author.get('scenes',[]))
    for chapter,ident in drafts.get('writing.sceneIds',{}).items():
        if not any(s['id']==ident for s in scenes):scenes.append({'id':ident,'chapter':int(chapter),'title':'Chapter '+chapter,'content':''})
    if scene_id:
        scenes=[s for s in scenes if s['id']==scene_id]
        if not scenes:raise ValueError('Choose a chapter from this project to download')
    rows=[]
    for scene in sorted(scenes,key=lambda s:s['chapter']):
        out=scene_output(folder,scene,drafts,live=False)
        if out['text'].strip():rows.append({'title':out['title'],'text':out['text'],'chapter':scene['chapter']})
    if not rows:raise ValueError('Write or generate a chapter before downloading')
    return author.get('title','Your writing'),contract({'medium':author.get('brief',{}).get('medium','Novel')}),rows


def blocks(text,script=False):
    """Recognize script cues without modifying the author's words."""
    dialogue=False
    for line in text.splitlines():
        if not line.strip():dialogue=False;yield 'blank','';continue
        stripped=line.strip()
        if script and re.match(r'^(?:INT\.|EXT\.|INT\./EXT\.|I/E\.|ACT\b|SCENE\b|PAGE\s+\d|PANEL\s+\d|SEQUENCE\b)',stripped,re.I):dialogue=False;yield 'heading',line
        elif script and re.match(r'^(VIDEO|AUDIO|SOUND|SFX|MUSIC|NARRATION|INTERVIEW)\s*:',stripped,re.I):dialogue=False;yield 'cue',line
        elif script and stripped.isupper() and len(stripped)<45 and not re.search(r'[.!?]$',stripped):dialogue=True;yield 'speaker',line
        elif script and dialogue:yield 'dialogue',line
        else:yield 'body',line


def render_docx(title,spec,rows):
    try:
        from docx import Document
        from docx.shared import Inches,Pt,RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError:raise ValueError('Install the studio export dependencies to download Word files') from None
    doc=Document();section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(1)
    section.left_margin=Inches(1.5 if spec['script'] else 1);section.right_margin=Inches(1)
    normal=doc.styles['Normal'];normal.font.name='Courier New' if spec['script'] else 'Times New Roman';normal.font.size=Pt(12);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=1 if spec['script'] else 1.5
    normal.paragraph_format.space_after=Pt(0 if spec['script'] else 6)
    for name in ('Normal','Title','Heading 1','Heading 2','Title Char','Heading 1 Char','Heading 2 Char'):
        if name not in doc.styles:continue
        style=doc.styles[name];style.font.name=normal.font.name;style.font.color.rgb=RGBColor(0,0,0)
        fonts=style.element.rPr.rFonts
        for attribute in list(fonts.attrib):
            if 'theme' in attribute.lower():del fonts.attrib[attribute]
        for border in style.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    doc.styles['Title'].font.size=Pt(18 if spec['script'] else 20)
    doc.styles['Heading 1'].font.size=Pt(14)
    doc.add_heading(title,0)
    doc.add_paragraph(spec['label'].capitalize() + ' · ' + str(sum(len(row['text'].split()) for row in rows)) + ' words')
    if len(rows)>1:
        doc.add_heading('Contents',1)
        for row in rows:doc.add_paragraph(row['title'])
    doc.add_page_break()
    for index,row in enumerate(rows):
        if index:doc.add_page_break()
        doc.add_heading(row['title'],1)
        for kind,text in blocks(row['text'],spec['script']):
            p=doc.add_paragraph(text)
            if kind=='heading':p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
            if kind=='speaker':p.paragraph_format.left_indent=Inches(2.2);p.paragraph_format.keep_with_next=True
            elif kind=='dialogue':p.paragraph_format.left_indent=Inches(1);p.paragraph_format.right_indent=Inches(1.5)
            elif kind=='cue':p.runs[0].bold=True
    footer=section.footer.paragraphs[0];footer.alignment=2
    section.header.paragraphs[0].text=title
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title=title;doc.core_properties.subject=spec['label']
    output=io.BytesIO();doc.save(output);return output.getvalue()


def render_pdf(title,spec,rows):
    try:
        from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,PageBreak
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.enums import TA_LEFT
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from xml.sax.saxutils import escape
    except ImportError:raise ValueError('Install the studio export dependencies to download PDF files') from None
    font='Courier' if spec['script'] else 'Times-Roman'
    choices=[Path('C:/Windows/Fonts')/('cour.ttf' if spec['script'] else 'times.ttf'),Path('/usr/share/fonts/truetype/dejavu')/('DejaVuSansMono.ttf' if spec['script'] else 'DejaVuSerif.ttf')]
    for path in choices:
        if path.is_file():
            font='ParaMythScript' if spec['script'] else 'ParaMythBook'
            if font not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont(font,str(path)))
            break
    output=io.BytesIO();doc=SimpleDocTemplate(output,pagesize=(612,792),leftMargin=108 if spec['script'] else 72,rightMargin=72,topMargin=72,bottomMargin=72,title=title)
    styles={k:ParagraphStyle(k,fontName=font,fontSize=12,leading=14 if spec['script'] else 18,spaceAfter=0 if spec['script'] else 6,alignment=TA_LEFT,splitLongWords=True) for k in ('body','heading','speaker','dialogue','cue')}
    styles['heading'].keepWithNext=True;styles['heading'].spaceBefore=10
    styles['speaker'].leftIndent=158;styles['speaker'].keepWithNext=True
    styles['dialogue'].leftIndent=72;styles['dialogue'].rightIndent=108
    title_style=ParagraphStyle('ProjectTitle',parent=styles['body'],fontSize=20,leading=26,spaceAfter=24)
    unit_style=ParagraphStyle('UnitTitle',parent=styles['body'],fontSize=15,leading=20,spaceAfter=18,keepWithNext=True)
    from ..document_exports import unicode_text
    elements=[Paragraph(unicode_text(escape(title)),title_style),Paragraph(escape(spec['label'].capitalize())+' · '+str(sum(len(row['text'].split()) for row in rows))+' words',styles['body'])]
    if len(rows)>1:
        elements.append(Paragraph('Contents',unit_style))
        elements.extend(Paragraph(escape(row['title']),styles['body']) for row in rows)
    elements.append(PageBreak())
    for index,row in enumerate(rows):
        if index:elements.append(PageBreak())
        elements.append(Paragraph(unicode_text(escape(row['title'])),unit_style))
        for kind,text in blocks(row['text'],spec['script']):
            elements.append(Spacer(1,12) if kind=='blank' else Paragraph(unicode_text(escape(text)),styles[kind]))
    def footer(canvas,doc):
        from ..document_exports import document_font
        canvas.setFont(document_font() or font,9);canvas.drawRightString(540,36,str(doc.page));canvas.drawString(doc.leftMargin,756,title[:80])
    doc.build(elements,onFirstPage=footer,onLaterPages=footer);return output.getvalue()


def export_writing(folder,kind,scene_id='',packet='manuscript'):
    if packet not in ('manuscript','project'):raise ValueError('Choose a manuscript or project dossier')
    try:title,spec,rows=manuscript(Path(folder),scene_id)
    except ValueError as error:
        if packet!='project' or str(error)!='Write or generate a chapter before downloading':raise
        current=Path(folder)
        author=load(current/'author-workspace.json') if (current/'author-workspace.json').is_file() else {}
        title=author.get('title','Your project');spec=contract({'medium':author.get('brief',{}).get('medium','Novel')});rows=[]
    if packet=='project':
        from ..document_exports import export_document,project_sections
        current=Path(folder)
        author=load(current/'author-workspace.json') if (current/'author-workspace.json').is_file() else {}
        reports=sorted((current/'research').glob('*.json'),key=lambda path:path.stat().st_mtime,reverse=True)[:1]
        sections=[{'title':row['title'],'paragraphs':[row['text']],'new_page':True,'script':spec['script']} for row in rows]
        body,mime=export_document(title,spec['label'].capitalize()+' project dossier',sections+project_sections(author,[load(path) for path in reports]),kind)
        stem=re.sub(r'[^\w -]','',title).strip()[:80] or 'project'
        return body,mime,stem+'-dossier.'+kind
    if kind=='docx':body=render_docx(title,spec,rows);mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    elif kind=='pdf':body=render_pdf(title,spec,rows);mime='application/pdf'
    elif kind=='fountain':
        if spec['medium'] not in ('screenplay','television'):raise ValueError('Fountain downloads are for movie and television scripts')
        body=('Title: '+title+'\n\n'+'\n\n'.join('# '+r['title']+'\n\n'+r['text'] for r in rows)).encode('utf-8');mime='text/plain; charset=utf-8'
    elif kind=='txt':body=(title+'\n\n'+'\n\n'.join(r['title']+'\n\n'+r['text'] for r in rows)).encode('utf-8');mime='text/plain; charset=utf-8'
    elif kind=='md':body=('# '+title+'\n\n'+'\n\n'.join('## '+r['title']+'\n\n'+r['text'] for r in rows)).encode('utf-8');mime='text/markdown; charset=utf-8'
    else:raise ValueError('Choose Word, PDF, Fountain or text')
    stem=re.sub(r'[^\w -]','',rows[0]['title'] if scene_id else title).strip()[:80] or 'writing'
    return body,mime,stem+'.'+kind
