"""Readable, paginated Word and PDF dossiers from structured project information."""
import io
import re
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr


def text(value):
    if isinstance(value, list):
        return '; '.join(text(item) for item in value)
    if isinstance(value, dict):
        return '; '.join(str(key).replace('_', ' ') + ': ' + text(item) for key, item in value.items() if item not in ('', None, [], {}))
    return str(value) if value is not None else ''


def readable_date(value):
    try:return datetime.fromisoformat(value.replace('Z','+00:00')).astimezone(timezone.utc).strftime('%d %B %Y, %H:%M UTC')
    except (ValueError,TypeError,AttributeError):return str(value)


def document_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    path=Path(__file__).parent/'assets/fonts/document-regular.ttf'
    if path.is_file():
        if 'ParaMythDocument' not in pdfmetrics.getRegisteredFontNames():pdfmetrics.registerFont(TTFont('ParaMythDocument',str(path)))
        return 'ParaMythDocument'
    return ''


def unicode_text(value):
    font=document_font()
    if not font:return value
    return ''.join(part if part.startswith('<') else re.sub(r'[^\x00-\x7f]+',lambda match:'<font name="'+font+'">'+match.group()+'</font>',part) for part in re.split(r'(<[^>]*>)',value))


def linked_text(value):
    parts, cursor = [], 0
    for match in re.finditer(r'https?://[^\s<>]+', text(value)):
        parts.append(escape(text(value)[cursor:match.start()]))
        url = match.group()
        parts.append('<link href=' + quoteattr(url) + ' color="#24416c">' + escape(url) + '</link>')
        cursor = match.end()
    parts.append(escape(text(value)[cursor:]))
    return unicode_text(''.join(parts).replace('\n', '<br/>'))


def research_sections(report):
    outline = report.get('outline', {})
    cfg = report['config']
    brief = [label + ': ' + cfg[key] for key, label in [('topic', 'Topic'), ('place', 'Place'), ('period', 'Period'), ('questions', 'Research questions')] if cfg.get(key)]
    sections = [{'title': 'Research brief', 'paragraphs': [report['notice'], *brief,
                 'Retrieved ' + readable_date(report['created_at']), *report.get('warnings', [])]}]
    if outline.get('angle'):
        sections.append({'title': 'Proposed approach', 'paragraphs': [outline['angle']]})
    if outline.get('claims'):
        sections.append({'title': 'Claims to verify', 'paragraphs': [str(index + 1) + '. ' + row['text'] + '\nSource leads: ' + ', '.join(row['source_ids']) + '\nVerification task: ' + row['check'] for index, row in enumerate(outline['claims'])]})
    if outline.get('timeline'):
        sections.append({'title': 'Provisional timeline', 'columns': ['Period', 'Event to verify', 'Source leads'], 'widths':[.18,.62,.2], 'rows': [[row['period'], row['event'], ', '.join(row['source_ids'])] for row in outline['timeline']]})
    for row in outline.get('interviews', []):
        sections.append({'title': 'Interview proposal ' + row['role'], 'paragraphs': [row['why'], *[str(index + 1) + '. ' + question for index, question in enumerate(row['questions'])]]})
    for row in outline.get('sequences', []):
        sections.append({'title': row['title'], 'paragraphs': ['Proposed video: ' + row['video'], 'Proposed audio: ' + row['audio'], 'Source leads: ' + (', '.join(row.get('source_ids', [])) or 'Creative proposal; research needed')]})
    if outline.get('archive_leads'):
        sections.append({'title': 'Archive research', 'paragraphs': [row['task'] + '\nSource leads: ' + ', '.join(row.get('source_ids', [])) for row in outline['archive_leads']]})
    for key, label in [('open_questions', 'Open questions'), ('care', 'Consent access and representation')]:
        if outline.get(key):sections.append({'title': label, 'paragraphs': [str(index + 1) + '. ' + row for index, row in enumerate(outline[key])]})
    sections.append({'title': 'Source register', 'paragraphs': [source['id'] + '  ' + source['title'] + '\n' + source['url'] +
                     '\nSearch excerpt: ' + source.get('snippet', '') + '\nVerification: ' + source['verification'] for source in report['sources']]})
    return sections


def production_sections(plan):
    cfg = plan.get('config', {})
    settings = [label + ': ' + text(cfg[key]) for key, label in [('region', 'Proposed filming region'), ('start_date', 'Proposed start date'), ('timezone', 'Schedule time zone'), ('currency', 'Budget currency'), ('crew_size', 'Crew assumption'), ('day_hours', 'Workday hours')] if cfg.get(key)]
    sections = [{'title': 'Planning summary', 'paragraphs': [plan.get('summary') or plan.get('recommendations', {}).get('summary', ''),
                 'Draft planning estimates. Confirm access, quotes, safety, dates and permissions before committing.',
                 'Prepared ' + readable_date(plan.get('created_at', '')), *settings]}]
    rec = plan.get('recommendations', {})
    for label, rows in [('Priority actions', rec.get('next_steps', [])), ('Location recommendations', rec.get('locations', [])),
                        ('Crew and equipment', rec.get('crew_equipment', [])), ('Legal and logistics', rec.get('legal_logistics', [])),
                        ('Risks and assumptions', rec.get('risks', []) + plan.get('warnings', []) + plan.get('breakdown', {}).get('assumptions', []))]:
        if rows:
            sections.append({'title': label, 'paragraphs': [text(row) for row in rows]})
    events = plan.get('schedule', {}).get('events', [])
    sections.append({'title': 'Shoot schedule', 'columns': ['Scene', 'Start and end', 'Location', 'Cast'], 'widths':[.22,.26,.3,.22],
                     'rows': [[row.get('summary', row.get('scene_id', '')), row.get('start', '') + '\n' + row.get('end', ''), row.get('location', ''), text(row.get('cast', []))] for row in events]})
    costs = plan.get('budget', {})
    money = lambda value: str(costs.get('currency', cfg.get('currency', ''))) + ' ' + format(float(value or 0), ',.2f')
    budget_notes = [('Priced costs only including contingency: ' if costs.get('unknown_costs') else 'Estimated total including contingency: ') + money(costs.get('total_low')) + ' to ' + money(costs.get('total_high')),
                    'Contingency: ' + str(costs.get('contingency_percent', 0)) + '%',
                    'Budget cap: ' + (money(costs['cap']) if costs.get('cap') else 'Not set'),
                    'The high estimate exceeds the cap.' if costs.get('over_cap') else 'Confirm all unknown costs before deciding feasibility.',
                    *costs.get('assumptions', [])]
    if costs.get('unknown_costs'):budget_notes.append('Budget incomplete. These costs are excluded; the range does not establish affordability: ' + ', '.join(costs['unknown_costs']))
    sections.append({'title': 'Budget range', 'paragraphs': budget_notes,
                     'columns': ['Item', 'Quantity', 'Low', 'High', 'Basis'], 'widths':[.3,.12,.12,.12,.34],
                     'rows': [[row.get('description', ''), text(row.get('quantity')) + ' ' + row.get('unit', ''), text(row.get('low_total')), text(row.get('high_total')), row.get('basis', '')] for row in costs.get('items', [])]})
    operations=plan.get('operations', {})
    if operations.get('departments'):
        sections.append({'title':'Department coverage','columns':['Department','Scope','Reason'],'widths':[.26,.17,.57],
            'rows':[[row.get('label',row.get('id','')).replace('_',' '),row.get('status','').replace('_',' '),row.get('rationale','')] for row in operations['departments']]})
    if operations.get('tasks'):
        sections.append({'title':'Department tasks','records':[{'title':row.get('id','')+' · '+row.get('task',''), 'paragraphs':[
            row.get('department','')+' · Owner: '+row.get('owner',''),
            'Deliverable: '+row.get('deliverable',''),
            'Proposed timing: '+readable_date(row.get('start',''))+' to '+readable_date(row.get('end',''))+' · '+text(row.get('duration_hours'))+' hours',
            'Depends on: '+(text(row.get('dependencies',[])) or 'No preceding task')+' · Scenes: '+(text(row.get('scene_ids',[])) or 'Whole production')
        ]} for row in operations['tasks']]})
    if operations.get('crew'):
        sections.append({'title':'Crew roster and paid days','columns':['Role','People','Prep / shoot / post days','Daily estimate','Basis'],'widths':[.25,.1,.2,.2,.25],
            'rows':[[row.get('role',''),row.get('count',''),' / '.join(text(row.get(key,0)) for key in ('prep_days','shoot_days','post_days')),
                money(row.get('daily_low'))+' to '+money(row.get('daily_high')) if row.get('daily_high') else 'Quote needed',row.get('basis','')] for row in operations['crew']]})
    for label,key,title_key,fields in [
        ('Sets and continuity','sets','name',[('approach','Approach'),('build_and_strike','Build and strike'),('props_continuity','Props and continuity'),('power_access_safety','Power access and safety'),('scene_ids','Scenes')]),
        ('Equipment requirements','equipment','item',[('quantity','Quantity'),('rental_days','Rental days'),('purpose','Purpose'),('scene_ids','Scenes')])]:
        if operations.get(key):sections.append({'title':label,'records':[{'title':row.get(title_key,label),'paragraphs':[name+': '+text(row[field]) for field,name in fields if row.get(field) not in ('',None,[],{})]} for row in operations[key]]})
    if operations.get('constraints_conflicts') or operations.get('assumptions'):
        sections.append({'title':'Scope decisions and unresolved constraints','paragraphs':operations.get('constraints_conflicts',[])+operations.get('assumptions',[])})
    for key, values in plan.get('local_logistics', {}).items():
        if isinstance(values, list) and values:
            sections.append({'title': key.replace('_', ' ').capitalize() + ' research leads', 'records':[
                {'title':row.get('title',key), 'paragraphs':list(dict.fromkeys([name+': '+text(row[field]) for field,name in [('address','Address'),('snippet','Search excerpt'),('phone','Phone'),('rating','Search rating'),('price_level','Price indicator'),('url','Source')] if row.get(field) not in ('',None,[],{})]))}
                if isinstance(row,dict) else {'title':'Research lead','paragraphs':[text(row)]} for row in values]})
    evidence = plan.get('evidence', [])
    if evidence:
        sections.append({'title': 'Research sources', 'records':[{'title':row.get('id','')+' · '+row.get('title','Source'), 'paragraphs':[
            row.get('url',''),row.get('snippet',''),'Evidence: '+row.get('evidence_level','search lead').replace('_',' ')+' · Verify the original source before use.'
        ]} for row in evidence]})
    return sections


def project_sections(author, reports=()):
    brief = author.get('brief', {})
    fields = [('premise', 'Premise'), ('medium', 'Writing format'), ('genre', 'Genre'), ('tone', 'Tone'), ('pov', 'Point of view'), ('themes', 'Themes'), ('worldRules', 'World rules'), ('mustAvoid', 'Constraints'), ('chapters', 'Planned chapters'), ('wordMin', 'Minimum scene words'), ('wordMax', 'Maximum scene words')]
    sections = [{'title': 'Project brief', 'paragraphs': [label + ': ' + text(brief[key]) for key, label in fields if brief.get(key)]}]
    for character in brief.get('characters', []):
        sections.append({'title': character.get('name') or 'Character', 'paragraphs': [label + ': ' + text(character[key]) for key, label in [('goal', 'Goal'), ('fear', 'Fear'), ('voice', 'Voice'), ('arc', 'Character arc')] if character.get(key)]})
    if author.get('canon'):
        sections.append({'title': 'Story canon', 'paragraphs': [text({'kind': row.get('kind'), 'text': row.get('text'), 'state': row.get('state')}) for row in author['canon']]})
    for report in reports:
        sections.extend(research_sections(report))
    return sections


def export_document(title, subtitle, sections, kind):
    if kind == 'docx':
        return docx_document(title, subtitle, sections), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if kind == 'pdf':
        return pdf_document(title, subtitle, sections), 'application/pdf'
    raise ValueError('Choose PDF or Word for this document')


def docx_document(title, subtitle, sections):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    page = doc.sections[0]
    page.top_margin = page.bottom_margin = Inches(.8)
    page.left_margin = page.right_margin = Inches(.85)
    page.page_width, page.page_height = Inches(8.5), Inches(11)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(9)
        if name == 'Normal':
            style.font.size = Pt(11); style.paragraph_format.line_spacing = 1.2
    doc.styles['Title'].font.size = Pt(28)
    doc.styles['Heading 1'].font.size = Pt(17)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    doc.add_paragraph(title, 'Title')
    doc.add_paragraph(subtitle, 'Subtitle')
    doc.add_paragraph('Contents', 'Heading 1')
    for index, section in enumerate(sections):
        para = doc.add_paragraph()
        link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'section_' + str(index))
        run = OxmlElement('w:r'); word = OxmlElement('w:t'); word.text = section['title']; run.append(word); link.append(run); para._p.append(link)
    doc.add_page_break()
    for index, section in enumerate(sections):
        if section.get('new_page') and index:
            doc.add_page_break()
        heading = doc.add_paragraph(section['title'], 'Heading 1')
        bookmark = OxmlElement('w:bookmarkStart'); bookmark.set(qn('w:id'), str(index)); bookmark.set(qn('w:name'), 'section_' + str(index)); heading._p.insert(0, bookmark)
        end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(index)); heading._p.append(end)
        content=list(section.get('paragraphs', []))
        for record in section.get('records',[]):content.extend([{'heading':record['title']},'\n'.join(str(value) for value in record.get('paragraphs',[]) if value)])
        for paragraph in content:
            if isinstance(paragraph,dict):
                doc.add_paragraph(paragraph['heading'],'Heading 2')
                continue
            if paragraph:
                if section.get('script'):
                    from .pipeline.writing_exports import blocks
                    for kind, line in blocks(str(paragraph), True):
                        para = doc.add_paragraph(line)
                        para.paragraph_format.space_after = Pt(0)
                        for run in para.runs:run.font.name = 'Courier New'; run.font.size = Pt(12)
                        if kind in ('heading', 'speaker'):para.paragraph_format.keep_with_next = True
                        if kind == 'heading':
                            for run in para.runs:run.bold = True
                        if kind == 'speaker':para.paragraph_format.left_indent = Inches(2)
                        if kind == 'dialogue':para.paragraph_format.left_indent = Inches(1); para.paragraph_format.right_indent = Inches(1)
                else:
                    from docx.opc.constants import RELATIONSHIP_TYPE
                    para, cursor = doc.add_paragraph(), 0
                    if '\n' in str(paragraph) and len(str(paragraph))<2500:para.paragraph_format.keep_together=True
                    for match in re.finditer(r'https?://[^\s<>]+', str(paragraph)):
                        para.add_run(str(paragraph)[cursor:match.start()])
                        link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), doc.part.relate_to(match.group(), RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
                        run = OxmlElement('w:r'); properties = OxmlElement('w:rPr'); colour = OxmlElement('w:color'); colour.set(qn('w:val'), '24416C'); properties.append(colour); run.append(properties)
                        word = OxmlElement('w:t'); word.text = match.group(); run.append(word); link.append(run); para._p.append(link)
                        cursor = match.end()
                    para.add_run(str(paragraph)[cursor:])
        if section.get('columns') and section.get('rows'):
            table = doc.add_table(rows=1, cols=len(section['columns'])); table.style = 'Table Grid'
            table.autofit=False
            widths=section.get('widths',[1/len(section['columns'])]*len(section['columns']))
            for column,width in zip(table.columns,widths):column.width=Inches(6.8*width)
            borders = OxmlElement('w:tblBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                border = OxmlElement('w:' + edge); border.set(qn('w:val'), 'single'); border.set(qn('w:sz'), '4'); border.set(qn('w:color'), 'D9D9D9'); borders.append(border)
            table._tbl.tblPr.append(borders)
            for cell, name, width in zip(table.rows[0].cells, section['columns'], widths):
                cell.width=Inches(6.8*width)
                cell.text = name
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EDF3'); cell._tc.get_or_add_tcPr().append(shade)
            repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in section['rows']:
                table_row=table.add_row()
                table_row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell, value, width in zip(table_row.cells, row, widths):
                    cell.width=Inches(6.8*width)
                    cell.text = text(value)
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(6)
                        for run in paragraph.runs:run.font.size=Pt(10)
            doc.add_paragraph()
    header = page.header.paragraphs[0]; header.text = title
    for run in header.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
    footer = page.footer.paragraphs[0]; footer.alignment = 2
    footer.add_run('Page ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    footer.add_run(' of ')
    total = OxmlElement('w:fldSimple'); total.set(qn('w:instr'), 'NUMPAGES'); footer._p.append(total)
    doc.core_properties.title, doc.core_properties.subject = title, subtitle
    output = io.BytesIO(); doc.save(output); return output.getvalue()


def pdf_document(title, subtitle, sections):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, LongTable, TableStyle, KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    font = document_font() or 'Helvetica'
    for path in (Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'), Path('C:/Windows/Fonts/arial.ttf')):
        if font == 'Helvetica' and path.is_file():
            font = 'ParaMythDossier'
            if font not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(font, str(path)))
            break
    body = ParagraphStyle('DossierBody', fontName=font, fontSize=10, leading=15, spaceAfter=9, splitLongWords=True)
    heading = ParagraphStyle('DossierHeading', parent=body, fontSize=17, leading=23, spaceBefore=18, spaceAfter=12, keepWithNext=True)
    large = ParagraphStyle('DossierTitle', parent=heading, fontSize=28, leading=34, spaceAfter=18)
    cell = ParagraphStyle('DossierCell', parent=body, fontSize=9, leading=13, spaceAfter=0)
    subheading=ParagraphStyle('DossierSubheading',parent=heading,fontSize=12,leading=17,spaceBefore=13,spaceAfter=7)
    script_styles = {kind: ParagraphStyle('DossierScript' + kind, parent=body, fontName='Courier', fontSize=12, leading=14, spaceAfter=0) for kind in ('body', 'heading', 'speaker', 'dialogue', 'cue')}
    script_styles['speaker'].leftIndent = 144; script_styles['speaker'].keepWithNext = True
    script_styles['dialogue'].leftIndent = 72; script_styles['dialogue'].rightIndent = 72
    script_styles['heading'].keepWithNext = True
    p = lambda value, style=body: Paragraph(linked_text(value), style)
    flow = [p(title, large), p(subtitle), p('Contents', heading)]
    for index, section in enumerate(sections):
        flow.append(Paragraph('<link href="#section_' + str(index) + '">' + escape(section['title']) + '</link>', body))
    flow.append(PageBreak())
    class Dossier(SimpleDocTemplate):
        def afterFlowable(self, item):
            if hasattr(item, 'section_index'):
                key = 'section_' + str(item.section_index)
                self.canv.bookmarkPage(key); self.canv.addOutlineEntry(item.getPlainText(), key)
    for index, section in enumerate(sections):
        if section.get('new_page') and index:
            flow.append(PageBreak())
        title_item = p(section['title'], heading); title_item.section_index = index; flow.append(title_item)
        content=list(section.get('paragraphs', []))
        for record in section.get('records',[]):content.extend([{'heading':record['title']},'\n'.join(str(value) for value in record.get('paragraphs',[]) if value)])
        for value in content:
            if isinstance(value,dict):
                flow.append(p(value['heading'],subheading))
                continue
            if section.get('script'):
                from .pipeline.writing_exports import blocks
                flow.extend(Spacer(1, 12) if kind == 'blank' else p(line, script_styles[kind]) for kind, line in blocks(value, True))
            elif value:
                item=p(value)
                flow.append(KeepTogether([item]) if '\n' in str(value) and len(str(value))<2500 else item)
        if section.get('columns') and section.get('rows'):
            rows = [[p(value, cell) for value in row] for row in [section['columns'], *section['rows']]]
            table = LongTable(rows, colWidths=[489*width for width in section.get('widths',[1/len(section['columns'])]*len(section['columns']))], repeatRows=1, splitInRow=1)
            table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E8EDF3')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F6F7F9')]),
                ('GRID', (0, 0), (-1, -1), .5, colors.HexColor('#D9D9D9')), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 7), ('RIGHTPADDING', (0, 0), (-1, -1), 7),
                ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
            flow.extend([table, Spacer(1, 12)])
    def furniture(canvas, document):
        canvas.setFont(font, 8); canvas.setFillColor(colors.black)
        canvas.drawString(61, 754, title[:85]); canvas.drawRightString(551, 35, str(document.page))
    output = io.BytesIO()
    Dossier(output, pagesize=(612, 792), leftMargin=61, rightMargin=62, topMargin=61, bottomMargin=60, title=title, author='ParaMyth').build(flow, onFirstPage=furniture, onLaterPages=furniture)
    return output.getvalue()
