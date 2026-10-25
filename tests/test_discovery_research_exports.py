"""Community isolation, cited documentary preparation and usable document exports."""
import io
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import Mock, patch
from uuid import uuid4
from zipfile import ZipFile

from fastapi.testclient import TestClient
from paramyth.studio.app import create_app
from paramyth.studio import research
from paramyth.storage import save
from paramyth.providers import ProviderError
from paramyth.document_exports import export_document, research_sections, production_sections
from paramyth.pipeline.writing_exports import export_writing
from test_studio_workspace import author_project


class DiscoveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app(self.temp.name, frontend=None)
        self.client = TestClient(self.app)
        self.headers = {'X-ParaMyth-Token': self.app.state.runtime.token}
        self.post = {'kind': 'story', 'title': 'A shared excerpt', 'body': 'This is the text I chose to share.', 'tags': ['Fantasy'],
                     'submission_id': str(uuid4()), 'confirm_public': True}

    def tearDown(self):
        self.app.state.runtime.close(); self.temp.cleanup()

    def publish(self, author='a'):
        with patch('paramyth.studio.discovery.actor', return_value=(author, 'Creator ' + author)):
            response = self.client.post('/api/ui/discovery', json=self.post, headers=self.headers)
        self.assertEqual(response.status_code, 201, response.text)
        return response.json()

    def test_private_projects_do_not_appear_and_posts_do_not_expose_private_ids(self):
        self.app.state.runtime.create_project('Private title', 'PRIVATE SOURCE NEVER PUBLIC')
        self.assertEqual(self.client.get('/api/ui/discovery').json()['posts'], [])
        post = self.publish()
        self.assertNotIn('author_id', post); self.assertNotIn('project_id', post)
        self.assertNotIn('PRIVATE SOURCE', json.dumps(post))

    def test_publication_needs_csrf_and_explicit_consent(self):
        self.assertEqual(self.client.post('/api/ui/discovery', json=self.post).status_code, 403)
        self.assertEqual(self.client.post('/api/ui/discovery', json={**self.post, 'confirm_public': False}, headers=self.headers).status_code, 422)

    def test_retried_submission_is_idempotent_and_survives_reopen(self):
        first = self.publish(); second = self.publish()
        self.assertEqual(first['id'], second['id'])
        app = create_app(self.temp.name, frontend=None)
        try:self.assertEqual(len(TestClient(app).get('/api/ui/discovery').json()['posts']), 1)
        finally:app.state.runtime.close()

    def test_ownership_revision_and_withdrawal(self):
        post = self.publish()
        edit = {**post, 'body': 'Updated deliberately shared writing.', 'confirm_public': True}
        with patch('paramyth.studio.discovery.actor', return_value=('b', 'Other creator')):
            self.assertEqual(self.client.patch('/api/ui/discovery/' + post['id'], json=edit, headers=self.headers).status_code, 409)
        with patch('paramyth.studio.discovery.actor', return_value=('a', 'Creator a')):
            changed = self.client.patch('/api/ui/discovery/' + post['id'], json=edit, headers=self.headers)
            self.assertEqual(changed.status_code, 200)
            self.assertEqual(self.client.patch('/api/ui/discovery/' + post['id'], json=edit, headers=self.headers).status_code, 409)
            withdrawn = self.client.patch('/api/ui/discovery/' + post['id'], json={**changed.json(), 'published': False}, headers=self.headers)
            self.assertEqual(withdrawn.status_code, 200)
            self.assertFalse(self.client.get('/api/ui/discovery?mine=true').json()['posts'][0]['published'])
        self.assertEqual(self.client.get('/api/ui/discovery').json()['posts'], [])
        self.assertEqual(self.client.get('/api/ui/discovery/' + post['id']).status_code, 404)

    def test_search_filters_and_reporting(self):
        post = self.publish()
        self.assertEqual(len(self.client.get('/api/ui/discovery?q=excerpt&kind=story').json()['posts']), 1)
        self.assertEqual(self.client.get('/api/ui/discovery?q=%25').json()['posts'], [])
        self.assertEqual(self.client.get('/api/ui/discovery?kind=character').json()['posts'], [])
        with patch('paramyth.studio.discovery.actor', return_value=('b', 'Creator b')):
            response = self.client.post('/api/ui/discovery/' + post['id'] + '/report', json={'reason': 'Spam'}, headers=self.headers)
            self.assertEqual(response.status_code, 200)
            self.assertEqual(self.client.get('/api/ui/discovery').json()['posts'], [])
        self.assertEqual(len(self.client.get('/api/ui/discovery').json()['posts']), 1)

    def test_post_validation_and_nonexistent_cursor(self):
        for update in ({'kind': 'other'}, {'body': ''}, {'tags': ['x'] * 6}, {'submission_id': '../private'}):
            self.assertEqual(self.client.post('/api/ui/discovery', json={**self.post, **update}, headers=self.headers).status_code, 422)
        self.assertEqual(self.client.get('/api/ui/discovery?before=missing').status_code, 400)


def lead():
    return {'title': 'Urban rivers original records', 'url': 'https://archive.example.org/records', 'snippet': 'A source excerpt, not verified full text.'}


class ResearchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.root = Path(self.temp.name)
        self.config = research.validate({'topic': 'Urban rivers', 'place': 'Seoul', 'period': '1950 to present', 'questions': 'Whose accounts disagree?'})

    def tearDown(self):self.temp.cleanup()

    def search(self):
        search = Mock()
        search.search.side_effect = lambda query, purpose: {'query': query, 'cached': False, 'retrieved_at': 100, 'results': [lead()]}
        return search

    def test_queries_follow_entered_topic_place_period_without_fixed_locations(self):
        queries = research.queries(self.config)
        self.assertEqual(len(queries), 4)
        self.assertTrue(all('Seoul' in query and 'Urban rivers' in query for query in queries))
        self.assertFalse(any('Morocco' in query or 'Ouarzazate' in query for query in queries))

    def test_synthesis_cannot_cite_invented_sources(self):
        with self.assertRaisesRegex(ValueError, 'supplied source'):
            research.validate_outline({'angle': 'An approach', 'claims': [{'text': 'Unverified claim', 'check': 'Read record', 'source_ids': ['S999']}]}, [{'id': 'S1'}])
        with self.assertRaisesRegex(ValueError, 'supporting source'):
            research.validate_outline({'angle': 'An approach', 'claims': [{'text': 'Unverified claim', 'check': 'Read record', 'source_ids': []}]}, [{'id': 'S1'}])

    def test_year_only_and_unrelated_search_results_are_excluded(self):
        self.assertFalse(research.relevant({'title': '1950 Census records', 'snippet': 'Population enumeration'}, 'Cheonggyecheon river restoration in Seoul'))
        self.assertTrue(research.relevant({'title': 'Cheonggyecheon Stream', 'snippet': 'The restoration project'}, 'Cheonggyecheon river restoration in Seoul'))

    def test_partial_search_failure_keeps_real_leads_without_fabricating_an_outline(self):
        search = self.search(); search.search.side_effect = [ProviderError('offline'), {'query': 'records', 'cached': True, 'retrieved_at': 100, 'results': [lead()]}, ProviderError('offline'), ProviderError('offline')]
        model = Mock(); model.ask.side_effect = ProviderError('offline')
        result = research.develop(self.root, self.config, lambda message: None, search=search, model=model)
        self.assertEqual(result['status'], 'leads_only'); self.assertEqual(result['outline'], {})
        self.assertEqual(len(result['sources']), 1); self.assertTrue(result['warnings'])
        self.assertTrue((self.root / 'research' / (result['id'] + '.json')).is_file())

    def test_search_only_does_not_call_model_and_deduplicates_sources(self):
        model = Mock()
        result = research.develop(self.root, {**self.config, 'synthesize': False}, lambda message: None, search=self.search(), model=model)
        model.ask.assert_not_called(); self.assertEqual(len(result['sources']), 1)
        self.assertEqual(result['sources'][0]['evidence_level'], 'search_snippet_only')

    def test_all_failed_searches_do_not_create_a_fake_success(self):
        search = Mock(); search.search.side_effect = ProviderError('offline')
        with self.assertRaises(ProviderError):research.develop(self.root, self.config, lambda message: None, search=search)
        self.assertFalse(list((self.root / 'research').glob('*.json')))

    def test_research_export_belongs_to_selected_project(self):
        app = create_app(self.root, frontend=None); runtime = app.state.runtime
        try:
            first = runtime.create_project('First', 'Private one'); second = runtime.create_project('Second', 'Private two')
            report = research.develop(runtime.folder(first), {**self.config, 'synthesize': False}, lambda message: None, search=self.search())
            client = TestClient(app)
            self.assertEqual(client.get('/api/ui/research?project_id=' + second).json()['reports'], [])
            self.assertEqual(client.get('/api/ui/research/' + report['id'] + '/export?project_id=' + second).status_code, 422)
            self.assertEqual(client.get('/api/ui/research/' + report['id'] + '/export?project_id=' + first).headers['content-type'], 'application/pdf')
        finally:runtime.close()


class DocumentTests(unittest.TestCase):
    def test_production_dossier_preserves_department_scope_tasks_and_crew(self):
        from paramyth.document_exports import production_sections
        from docx import Document
        from pypdf import PdfReader
        plan={'operations':{'departments':[{'label':'Archive rights','status':'planned','rationale':'Clear the ledger reproduction'}],
            'tasks':[{'id':'rights-1','task':'Request reproduction consent','owner':'producer','deliverable':'Written archive approval','department':'Archive rights','start':'2027-01-18T09:00:00+09:00','end':'2027-01-18T11:00:00+09:00','duration_hours':2}],
            'crew':[{'role':'Archive researcher','count':1,'prep_days':2,'shoot_days':1,'post_days':0,'daily_low':0,'daily_high':0,'basis':'Quote needed'}]}}
        sections=production_sections(plan)
        word,_=export_document('River production','Planning dossier',sections,'docx')
        doc=Document(io.BytesIO(word))
        content='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        pdf,_=export_document('River production','Planning dossier',sections,'pdf')
        extracted='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
        for value in ('Written archive approval','Clear the ledger reproduction','Archive researcher','Quote needed'):
            self.assertIn(value,content);self.assertIn(value,extracted)

    def test_scoped_script_revision_retries_context_leaks_and_preserves_surroundings(self):
        from paramyth.pipeline import creative_studio
        prefix='VIDEO:\nA careful observation of the river keeps the camera absolutely still for the first moment.\n\n'
        target='The camera holds for a long quiet beat.'
        suffix='\n\nAUDIO:\nOnly the water and a distant traffic hum are heard, without narration or music.'
        source=prefix+target+suffix
        captured=[]
        def complete(provider,model,messages,*args,**kwargs):
            captured.append(json.loads(messages[-1]['content']))
            replacement=prefix+'The camera holds for thirty seconds.' if len(captured)==1 else 'The camera holds for thirty uninterrupted seconds.'
            return {'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'title':'River','text':replacement,'assumptions':[]})}}]}
        with tempfile.TemporaryDirectory() as path,patch.object(creative_studio.creative_stream,'configuration',return_value=('test','test')),patch.object(creative_studio.creative_stream,'completion',side_effect=complete),patch.object(creative_studio.time,'sleep'):
            result=creative_studio.run(Path(path),{'operation':'revise','medium':'Documentary','text':source,'target_passage':target,'feedback':'Hold the moment longer','use_story':False},stream=True)
        self.assertEqual(len(captured),2)
        self.assertNotIn('target_words',captured[0])
        self.assertEqual(result['result']['text'],prefix+'The camera holds for thirty uninterrupted seconds.'+suffix)
        self.assertEqual(result['result']['text'].count('VIDEO:'),1)
        self.assertEqual(creative_studio.replacement_text(source,target,prefix+'A new beat.'+suffix),'A new beat.')
        with self.assertRaisesRegex(ValueError,'script or chapter headings'):
            creative_studio.replacement_text(source,target,'VIDEO:\nEXT. RIVER - DAWN\n\nThe camera holds for thirty seconds.')

    def test_pdf_embeds_multilingual_source_names(self):
        from pypdf import PdfReader
        body,_=export_document('Research dossier','Source register',[{'title':'Original source','paragraphs':['서울 청계천 · café · Αθήνα']}],'pdf')
        content='\n'.join(page.extract_text() for page in PdfReader(io.BytesIO(body)).pages)
        self.assertIn('서울 청계천',content)
        self.assertIn(b'/FontFile2',body)
    def test_project_dossier_works_before_a_chapter_is_written(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as path:
            root=Path(path);project=author_project('Planning stage');project['scenes']=[]
            save(root/'author-workspace.json',project)
            body,_,_=export_writing(root,'docx',packet='project')
            self.assertIn('Project brief','\n'.join(para.text for para in Document(io.BytesIO(body)).paragraphs))

    def test_local_crew_does_not_get_hotel_or_flight_allowances(self):
        from paramyth.pipeline.production import normalize_config,schedule
        from paramyth.pipeline.production_operations import research_logistics,excluded_for_local_crew
        from test_production import scene,breakdown
        cfg=normalize_config({'region':'Seoul','brief':'Use a three-person local crew','crew_size':3,'start_date':'2027-01-18'})
        self.assertTrue(cfg['local_crew'])
        search=Mock();search.search.return_value={'results':[]}
        source=breakdown();source['budget_items'].append({'description':'Crew hotel','category':'travel','quantity':3,'unit':'night','low_rate':100,'high_rate':200,'basis':'estimate'})
        result=research_logistics(search,cfg,schedule([scene()],cfg),[],[],[],source,{},lambda _:None)
        self.assertEqual(result['rooms'],0)
        self.assertTrue(all(call.kwargs['engine']=='google_maps' for call in search.search.call_args_list))
        self.assertFalse(any('hotel' in row['description'].lower() for row in source['budget_items']))
        self.assertFalse(excluded_for_local_crew({'category':'locations','description':'Hotel courtyard filming venue','unit':'day'}))

    def test_research_and_plan_json_round_trip_non_latin_place_names(self):
        from paramyth.storage import load
        with tempfile.TemporaryDirectory() as path:
            target=Path(path)/'plan.json'; value={'location':'서울 청계천','note':'A resident’s account'}
            save(target,value)
            with patch.object(Path,'read_text',autospec=True,wraps=Path.read_text) as read:
                # Verify the shared loader requests UTF-8 explicitly, even on a non-UTF-8 OS.
                read.return_value=json.dumps(value)
                self.assertEqual(load(target),value)
                self.assertEqual(read.call_args.kwargs['encoding'],'utf-8-sig')

    def test_production_repair_keeps_record_lists_structured_and_normalizes_task_ids(self):
        import copy
        from paramyth.pipeline import production
        from paramyth.pipeline.production_operations import develop_operations
        from test_advanced_workflow_depth import operations
        from test_production import scene, breakdown
        cfg=production.normalize_config({'start_date':'2027-01-18'})
        value=operations(); bad=copy.deepcopy(value)
        bad['departments'][0]['tasks'][0]['deliverable']=''
        model=Mock(); model.ask.side_effect=[bad,value]
        result=develop_operations(model,breakdown(),cfg,production.schedule([scene()],cfg),lambda _:None)
        self.assertEqual(len(result['tasks']),18)
        self.assertIn('must remain objects',model.ask.call_args.args[0])
        value=operations(); task=value['departments'][0]['tasks'][0]
        task['id']=42; task['action']=task.pop('task')
        value['departments'][3].update(status='not_applicable',tasks=[{}])
        model.ask.side_effect=None; model.ask.return_value=value
        result=develop_operations(model,breakdown(),cfg,production.schedule([scene()],cfg),lambda _:None)
        self.assertEqual(result['tasks'][0]['id'],'42')
        self.assertEqual(len(result['tasks']),17)

    def test_dossiers_have_contents_numbered_pages_sources_and_expandable_tables(self):
        from docx import Document
        from pypdf import PdfReader
        sections = [{'title': 'Source register', 'paragraphs': ['S1 Original record\nhttps://archive.example.org/records\nVerify the original source.']},
                    {'title': 'Detailed schedule', 'columns': ['Scene', 'Work', 'Owner'], 'rows': [[str(i), 'A detailed task with access and equipment requirements. ' * 4, 'Production'] for i in range(60)]}]
        docx, mime = export_document('Detailed production dossier', 'Draft plan', sections, 'docx')
        doc = Document(io.BytesIO(docx))
        self.assertEqual(len(doc.tables[0].rows), 61)
        with ZipFile(io.BytesIO(docx)) as archive:
            self.assertIn(b'NUMPAGES', archive.read('word/footer1.xml'))
            self.assertIn(b'bookmarkStart', archive.read('word/document.xml'))
        pdf, mime = export_document('Detailed production dossier', 'Draft plan', sections, 'pdf')
        reader = PdfReader(io.BytesIO(pdf))
        self.assertGreater(len(reader.pages), 3)
        self.assertIn('Verify the original source', ''.join(page.extract_text() for page in reader.pages))
        self.assertEqual(len(reader.outline), 2)

    def test_project_packet_contains_current_writing_characters_and_research(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as path:
            root = Path(path); project = author_project('A current project')
            project['scenes'][0]['content'] = 'The current author edit remains here.'
            save(root / 'author-workspace.json', project)
            report = research.develop(root, research.validate({'topic': 'Urban rivers', 'synthesize': False}), lambda message: None, search=ResearchTests.search(self))
            body, mime, name = export_writing(root, 'docx', packet='project')
            content = '\n'.join(para.text for para in Document(io.BytesIO(body)).paragraphs)
            self.assertIn('The current author edit remains here', content)
            self.assertIn('Project brief', content); self.assertIn('Source register', content)
            self.assertIn('archive.example.org/records', content)
            self.assertTrue(name.endswith('-dossier.docx'))

    def test_one_step_visual_job_preserves_source_and_only_renders_once(self):
        from paramyth.pipeline import image_studio as renderer
        with tempfile.TemporaryDirectory() as path:
            app = create_app(path, frontend=None); runtime = app.state.runtime
            try:
                ident = runtime.create_project('A visual project', 'A shore at dawn')
                prepared = {'prompt': 'A quiet shore', 'title': 'Dawn'}
                record = {'id': 'a' * 24, 'request': {'prompt': 'A quiet shore'}, 'images': [], 'status': 'completed'}
                with patch.object(renderer.CATALOG, 'get', return_value=renderer.MUSE_MODEL), patch.object(renderer, 'prepare_prompt', return_value=prepared) as prepare, patch.object(renderer, 'generate', return_value=record) as generate:
                    response = TestClient(app).post('/api/images/create?project_id=' + ident, headers={'X-ParaMyth-Token': runtime.token}, json={'model': 'muse-image-1.0', 'kind': 'scene'})
                    self.assertEqual(response.status_code, 202, response.text)
                    for attempt in range(100):
                        job = runtime.job(response.json()['job_id'])
                        if job['status'] not in ('queued', 'running'):break
                        time.sleep(.01)
                    self.assertEqual(job['status'], 'complete', job)
                    prepare.assert_called_once(); generate.assert_called_once()
                    self.assertIn('writing_signature', job['result'])
            finally:runtime.pool.shutdown(wait=True); runtime.close()
